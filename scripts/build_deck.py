"""Build the demo deck as a .pptx from docs/demo/deck.md.

The markdown stays the source of truth. It is the thing that gets edited, reviewed and
committed, and a slide file that is generated from it cannot drift away from the script
the way a hand-built one does. Re-run this after any edit to the deck:

    uv run --with python-pptx python scripts/build_deck.py

What lands on a slide is only the `**On screen**` block, which is deliberately thin. The
`**Script**` and `**Why this slide**` blocks go into the speaker notes, where they belong
and where the room cannot see them.

Where a slide calls for a recording or a screenshot that does not exist yet, the slide
gets a marked placeholder carrying the instruction verbatim, rather than being silently
dropped. A clip is picked up automatically once `docs/demo/clips/clip-N.mp4` exists.
"""
from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "docs" / "demo" / "deck.md"
CLIPS = ROOT / "docs" / "demo" / "clips"
OUT = ROOT / "docs" / "demo" / "keel-demo.pptx"

# The product's own palette, inverted. Several slides are full-bleed terminal video, and a
# light deck flashing to a dark recording and back is hard on a projected room.
INK = RGBColor(0x17, 0x20, 0x3A)        # navy-900, the background
PANEL = RGBColor(0x23, 0x2D, 0x4A)      # navy-800, a raised block
LINE = RGBColor(0x38, 0x41, 0x5F)       # navy-700, hairlines
MUTED = RGBColor(0x97, 0xA0, 0xB5)      # navy-400, secondary text
FAINT = RGBColor(0x63, 0x6E, 0x88)      # navy-500, slide numbers
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xD6, 0x29, 0x4E)     # crimson-600
AMBER = RGBColor(0xE0, 0xA0, 0x21)

SANS = "Segoe UI"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.9)


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
@dataclass
class Block:
    """One run of on-screen content: a quote, a code listing, bullets or a note."""

    kind: str          # quote | code | bullets | prose
    lines: list[str] = field(default_factory=list)


@dataclass
class Slide:
    number: str
    title: str
    onscreen_note: str = ""            # the text after "**On screen** - ..."
    blocks: list[Block] = field(default_factory=list)
    script: str = ""
    why: str = ""
    act: str = ""
    diagram: str = ""

    @property
    def is_video(self) -> bool:
        # Marked, not sniffed. Matching on the words "video" and "full bleed" meant the
        # deck could not say it in its own language without the clip becoming a
        # placeholder, which is the builder dictating prose to the writer.
        return self.diagram == "clip"


def _split_sections(text: str) -> list[tuple[str, str]]:
    """Return (heading, body) for every `#`/`##` heading, stopping at presenter material."""
    out, heading, buf = [], None, []
    for raw in text.splitlines():
        if raw.startswith("## Running order") or raw.startswith("## If you are running over"):
            break
        m = re.match(r"^(#{1,2}) (.+?)\s*$", raw)
        if m:
            if heading is not None:
                out.append((heading, "\n".join(buf)))
            heading, buf = raw, []
            continue
        buf.append(raw)
    if heading is not None:
        out.append((heading, "\n".join(buf)))
    return out


def _parse_onscreen(body: str) -> tuple[str, list[Block]]:
    """Pull the `**On screen**` section apart into a lead note and typed blocks."""
    m = re.search(r"\*\*On screen\*\*(.*?)(?=\n\*\*Script\*\*|\n\*\*Why this slide\*\*|\Z)",
                  body, re.S)
    if not m:
        return "", []
    chunk = m.group(1)

    first, _, rest = chunk.partition("\n")
    note = first.strip().lstrip("-").strip()

    blocks: list[Block] = []
    in_code = False
    for line in rest.splitlines():
        if line.startswith("```"):
            in_code = not in_code
            if in_code:
                blocks.append(Block("code"))
            continue
        if in_code:
            blocks[-1].lines.append(line)
            continue
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith(">"):
            text = stripped.lstrip("> ").rstrip()
            if blocks and blocks[-1].kind == "quote":
                blocks[-1].lines.append(text)
            else:
                blocks.append(Block("quote", [text]))
        elif stripped.startswith(("- ", "* ")):
            text = stripped[2:].strip()
            if blocks and blocks[-1].kind == "bullets":
                blocks[-1].lines.append(text)
            else:
                blocks.append(Block("bullets", [text]))
        else:
            blocks.append(Block("prose", [stripped]))
    return note, blocks


def _field(body: str, name: str) -> str:
    m = re.search(rf"\*\*{name}\*\*(.*?)(?=\n\*\*[A-Z]|\Z)", body, re.S)
    if not m:
        return ""
    text = m.group(1)
    text = re.sub(r"^\s*\(\d+:\d+\)", "", text).strip()
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)      # links to plain text
    return re.sub(r"\*\*|`", "", text).strip()


def parse(text: str) -> list[Slide]:
    slides: list[Slide] = []
    act = ""
    for heading, body in _split_sections(text):
        if heading.startswith("# Act"):
            act = heading[2:].strip()
            slides.append(Slide(number="", title=act, act=act,
                                script=body.strip(), blocks=[]))
            slides[-1].onscreen_note = "__ACT__"
            continue
        if not heading.startswith("## Slide"):
            continue
        m = re.match(r"## Slide (\S+)\s*-\s*(.+)", heading)
        if not m:
            continue
        note, blocks = _parse_onscreen(body)
        drawing = re.search(r"<!--\s*diagram:\s*(\w+)\s*-->", body)
        slides.append(Slide(
            number=m.group(1), title=m.group(2).strip(), act=act,
            onscreen_note=note, blocks=blocks, diagram=drawing.group(1) if drawing else "",
            script=_field(body, "Script"), why=_field(body, "Why this slide"),
        ))
    return slides


# --------------------------------------------------------------------------- #
# Drawing
# --------------------------------------------------------------------------- #
def _blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = INK
    return slide


def _text(slide, left, top, width, height, *, size, color=PAPER, bold=False,
          font=SANS, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bold, font, color
    return tf


def _para(tf, text, *, size, color=PAPER, bold=False, font=SANS,
          align=PP_ALIGN.LEFT, spacing=1.25, space_before=0):
    p = tf.add_paragraph() if tf.paragraphs[0].runs or len(tf.paragraphs) > 1 else tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    p.space_before = Pt(space_before)
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bold, font, color
    return p


def _fit(lines: list[str], width_in: float, height_in: float,
         ladder: tuple[int, ...], *, spacing: float = 1.3, mono: bool = False) -> int:
    """Largest point size from `ladder` at which the lines still fit the box.

    A PowerPoint text box does not clip: text that does not fit runs off the bottom of
    the slide and nothing in the file says so. Three slides shipped that way in the first
    build, so the size is chosen against the box rather than against the longest line.
    """
    per_char = 0.0088 if mono else 0.0075       # inches of width per point of size
    for size in ladder:
        used = 0.0
        for line in lines:
            wrapped = max(1, int(len(line) * per_char * size / max(width_in, 0.1)) + 1)
            used += wrapped * size * spacing / 72
        if used <= height_in:
            return size
    return ladder[-1]


def _rule(slide, left, top, width, color=ACCENT, thickness=Inches(0.035)):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, thickness)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _panel(slide, left, top, width, height, color=PANEL, outline=None):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.02
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if outline:
        shape.line.color.rgb = outline
        shape.line.width = Pt(1.25)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _notes(slide, s: Slide):
    parts = []
    if s.script:
        parts.append(s.script)
    if s.why:
        parts.append("--- why this slide ---\n" + s.why)
    slide.notes_slide.notes_text_frame.text = "\n\n".join(parts) or " "


def _footer(slide, s: Slide):
    if not s.number:
        return
    _text(slide, W - MARGIN - Inches(1.6), H - Inches(0.62), Inches(1.6), Inches(0.3),
          size=11, color=FAINT, align=PP_ALIGN.RIGHT)
    tf = slide.shapes[-1].text_frame
    tf.paragraphs[0].runs[0].text = s.number
    if s.act:
        _text(slide, MARGIN, H - Inches(0.62), Inches(6), Inches(0.3), size=11, color=FAINT)
        # The label only. An act's full name is a sentence and wraps into two lines here.
        slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = s.act.split(" - ")[0].upper()


# --------------------------------------------------------------------------- #
# Slide kinds
# --------------------------------------------------------------------------- #
def draw_title(prs, s: Slide):
    slide = _blank(prs)
    quote = next((b for b in s.blocks if b.kind == "quote"), Block("quote", ["Keel"]))
    lines = quote.lines
    _rule(slide, MARGIN, Inches(2.35), Inches(1.1))
    tf = _text(slide, MARGIN, Inches(2.7), W - 2 * MARGIN, H - Inches(3.5),
               size=66, bold=True, spacing=1.0)
    tf.paragraphs[0].runs[0].text = lines[0]
    if len(lines) > 1:
        _para(tf, lines[1], size=25, color=MUTED, spacing=1.3, space_before=18)
    if len(lines) > 2:
        _para(tf, lines[2], size=15, color=FAINT, font=MONO, space_before=26)
    _notes(slide, s)
    return slide


def draw_act(prs, s: Slide):
    slide = _blank(prs)
    label, _, name = s.title.partition(" - ")
    _text(slide, MARGIN, Inches(2.9), Inches(6), Inches(0.4), size=14, bold=True, color=ACCENT)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = label.upper()
    _text(slide, MARGIN, Inches(3.35), W - 2 * MARGIN, Inches(1.6), size=48, bold=True,
          spacing=1.1)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = name or label
    if s.script:
        _text(slide, MARGIN, Inches(4.95), Inches(8.2), Inches(1.0), size=16, color=MUTED,
              spacing=1.45)
        slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = s.script.split("\n")[0]
    slide.notes_slide.notes_text_frame.text = s.script or " "
    return slide


def draw_video(prs, s: Slide, ordinal: int):
    slide = _blank(prs)
    # Named by its position in the deck, not by slide number, because the shotlist calls
    # them clip 1, 2, 3. Keying on the slide number silently repointed every file the
    # first time a slide was inserted ahead of one.
    clip = CLIPS / f"clip-{ordinal}.mp4"
    left, top = Inches(0.55), Inches(0.75)
    width, height = W - 2 * left, H - 2 * top
    if clip.exists():
        poster = clip.with_suffix(".png")
        slide.shapes.add_movie(str(clip), left, top, width, height,
                               poster_frame_image=str(poster) if poster.exists() else None)
    else:
        _panel(slide, left, top, width, height, color=PANEL, outline=LINE)
        tf = _text(slide, left, top + height / 2 - Inches(0.75), width, Inches(1.5),
                   size=17, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tf.paragraphs[0].runs[0].text = s.title
        _para(tf, f"drop the recording at  docs/demo/clips/clip-{ordinal}.mp4  "
                  f"and rebuild", size=13, color=FAINT, font=MONO,
              align=PP_ALIGN.CENTER, space_before=14)
        _para(tf, s.onscreen_note, size=12, color=FAINT, align=PP_ALIGN.CENTER,
              space_before=10)
    _notes(slide, s)
    _footer(slide, s)
    return slide


def draw_statement(prs, s: Slide, quote: Block):
    slide = _blank(prs)
    _rule(slide, MARGIN, Inches(1.35), Inches(0.85))
    _text(slide, MARGIN, Inches(1.75), Inches(9), Inches(0.4), size=14, bold=True, color=FAINT)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = s.title.upper()

    top = Inches(2.55)
    box_h = H - top - Inches(1.0)
    width_in = (W - 2 * MARGIN) / 914400
    size = _fit(quote.lines, width_in, box_h / 914400,
                (44, 40, 36, 32, 28, 25, 22, 20, 18), spacing=1.35)
    tf = _text(slide, MARGIN, top, W - 2 * MARGIN, box_h,
               size=size, bold=True, spacing=1.22, anchor=MSO_ANCHOR.MIDDLE)
    tf.paragraphs[0].runs[0].text = quote.lines[0]
    for line in quote.lines[1:]:
        _para(tf, line, size=size, bold=True, spacing=1.22,
              space_before=min(10, size // 3))
    _notes(slide, s)
    _footer(slide, s)
    return slide


def _box(slide, left, top, width, height, title, caption="", *,
         fill=PANEL, outline=None, title_color=PAPER, title_size=13,
         title_font=SANS, caption_color=MUTED):
    shape = _panel(slide, left, top, width, height, color=fill, outline=outline)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.1
    run = p.add_run()
    run.text = title
    run.font.size, run.font.bold = Pt(title_size), True
    run.font.name, run.font.color.rgb = title_font, title_color
    if caption:
        _para(tf, caption, size=9, color=caption_color, align=PP_ALIGN.CENTER,
              spacing=1.15, space_before=3)
    return shape


def _arrow(slide, x1, y1, x2, y2, color=LINE, width=Pt(1.5)):
    """A straight connector with a head, which python-pptx does not expose."""
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml.ns import qn

    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = width
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return conn


def draw_spine(prs, s: Slide):
    """The one architecture slide, drawn rather than typed.

    Every box here is a real PowerPoint shape, so it survives being nudged on the night
    and it does not fall apart when the labels are translated, which is the whole reason
    the monospace version had to go.
    """
    slide = _blank(prs)
    _rule(slide, MARGIN, Inches(0.75), Inches(0.85))
    _text(slide, MARGIN, Inches(1.05), Inches(10.5), Inches(0.5), size=26, bold=True)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = s.title

    w, h = Inches(2.2), Inches(0.85)
    xs = [Inches(1.14 + 2.95 * i) for i in range(4)]
    spine_y = Inches(2.95)
    feed_y, feed_h = Inches(1.88), Inches(0.5)

    # What feeds the chain, sitting above the two boxes it feeds.
    _box(slide, xs[1], feed_y, w, feed_h, "surface", fill=INK, outline=LINE,
         title_size=12, title_color=MUTED)
    _box(slide, xs[2], feed_y, w, feed_h, "source", fill=INK, outline=LINE,
         title_size=12, title_color=MUTED)
    _arrow(slide, xs[1] + w // 2, feed_y + feed_h, xs[1] + w // 2, spine_y)
    _arrow(slide, xs[2] + w // 2, feed_y + feed_h, xs[2] + w // 2, spine_y)

    # Field names stay in English: they are what the room sees inside the tool during
    # the clips, and a translated label on the slide would not match the screen.
    spine = [
        ("component", "на чём она сидит", PANEL, None),
        ("weakness", "архитектурное условие", PANEL, None),
        ("threat", "опирается на одну или несколько", PANEL, ACCENT),
        ("harm", "ровно один", ACCENT, None),
    ]
    for x, (title, caption, fill, outline) in zip(xs, spine):
        # Muted grey is unreadable on the accent fill, which is the one box that has one.
        _box(slide, x, spine_y, w, h, title, caption, fill=fill, outline=outline,
             title_size=15,
             caption_color=RGBColor(0xF6, 0xC9, 0xD3) if fill == ACCENT else MUTED)
    for i in range(3):
        _arrow(slide, xs[i] + w, spine_y + h // 2, xs[i + 1], spine_y + h // 2)

    # Everything that hangs off the threat, grouped so it reads as belonging to it.
    gx, gy = Inches(2.55), Inches(4.35)
    gw, gh = Inches(9.6), Inches(1.4)
    _panel(slide, gx, gy, gw, gh, color=PANEL)
    _arrow(slide, xs[2] + w // 2, spine_y + h, xs[2] + w // 2, gy)

    pill_w, pill_h = Inches(2.85), Inches(0.5)
    pill_y = gy + Inches(0.2)
    hangs = [
        ("reachability", "когда здесь это вообще не рабочий путь", AMBER),
        ("mitigation · gating", "закрывает угрозу", RGBColor(0x1F, 0x9D, 0x57)),
        ("mitigation · soft", "снижает шанс, но не закрывает", FAINT),
    ]
    for i, (title, caption, colour) in enumerate(hangs):
        x = gx + Inches(0.28) + i * Inches(3.05)
        _box(slide, x, pill_y, pill_w, pill_h, title, fill=INK, outline=colour,
             title_size=12, title_color=colour)
        _text(slide, x, pill_y + pill_h + Inches(0.12), pill_w, Inches(0.5),
              size=10, color=MUTED, align=PP_ALIGN.CENTER, spacing=1.25)
        slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = caption

    # The bottom band: who runs it, through what, onto what.
    band_y, band_h = Inches(5.92), Inches(0.5)
    band = [("кто оценивает", Inches(1.14), Inches(2.6)),
            ("Keel · MCP + REST", Inches(4.5), Inches(2.9)),
            ("catalog/*.yaml в git", Inches(8.2), Inches(3.3))]
    for title, x, bw in band:
        mono = title.startswith("catalog")
        _box(slide, x, band_y, bw, band_h, title, fill=INK, outline=LINE,
             title_size=11, title_color=MUTED, title_font=MONO if mono else SANS)
    for x, bw, caption in ((Inches(1.14), Inches(2.6), "модель с навыком или человек"),
                           (Inches(8.2), Inches(3.3), "никакой базы, каждая правка это дифф файла")):
        _text(slide, x, band_y + band_h + Inches(0.09), bw, Inches(0.3),
              size=10, color=FAINT, align=PP_ALIGN.CENTER)
        slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = caption
    _arrow(slide, Inches(3.74), band_y + band_h // 2, Inches(4.5), band_y + band_h // 2)
    _arrow(slide, Inches(7.4), band_y + band_h // 2, Inches(8.2), band_y + band_h // 2)

    _notes(slide, s)
    _footer(slide, s)
    return slide


# Slides whose on-screen content is a drawing. The deck marks one with an HTML comment,
# `<!-- diagram: spine -->`, rather than being matched on its title: a title is prose and
# gets rewritten, and the first translation would have silently dropped the drawing.
DIAGRAMS = {"spine": draw_spine}


def draw_code(prs, s: Slide, code: Block):
    slide = _blank(prs)
    _rule(slide, MARGIN, Inches(0.75), Inches(0.85))
    _text(slide, MARGIN, Inches(1.1), Inches(10), Inches(0.5), size=24, bold=True)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = s.title

    body = list(code.lines)
    top = Inches(1.95)
    height = H - top - Inches(0.95)
    # Monospace art must not wrap, so the width ladder is the binding one; the height
    # check then stops a long listing from running off the bottom.
    longest = max((len(ln) for ln in body), default=40)
    by_width = 15 if longest <= 62 else 13 if longest <= 78 else 11 if longest <= 94 else 9
    by_height = _fit(body, (W - 2 * MARGIN - Inches(0.8)) / 914400,
                     (height - Inches(0.6)) / 914400,
                     (15, 13, 12, 11, 10, 9, 8), spacing=1.18, mono=True)
    size = min(by_width, by_height)
    _panel(slide, MARGIN, top, W - 2 * MARGIN, height, color=PANEL)
    tf = _text(slide, MARGIN + Inches(0.4), top + Inches(0.32),
               W - 2 * MARGIN - Inches(0.8), height - Inches(0.6),
               size=size, color=PAPER, font=MONO, spacing=1.18)
    tf.paragraphs[0].runs[0].text = body[0] if body else ""
    for line in body[1:]:
        _para(tf, line, size=size, font=MONO, spacing=1.18)
    _notes(slide, s)
    _footer(slide, s)
    return slide


def draw_bullets(prs, s: Slide, blocks: list[Block]):
    slide = _blank(prs)
    _rule(slide, MARGIN, Inches(0.75), Inches(0.85))
    _text(slide, MARGIN, Inches(1.1), Inches(10.5), Inches(0.6), size=32, bold=True)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = s.title

    top = Inches(2.3)
    box_h = H - top - Inches(0.95)
    lines = [ln for b in blocks for ln in b.lines]
    size = _fit(lines, (W - 2 * MARGIN - Inches(1.0)) / 914400, box_h / 914400,
                (24, 22, 20, 18, 16, 15, 14), spacing=1.5)
    tf = _text(slide, MARGIN, top, W - 2 * MARGIN - Inches(1.0), box_h,
               size=size, color=PAPER, spacing=1.35)
    started = False
    for b in blocks:
        for line in b.lines:
            mono = bool(re.fullmatch(r"[\d\s./%+-]+|[A-Z-]{3,}[\w\s-]*", line)) and len(line) < 24
            if not started:
                tf.paragraphs[0].runs[0].text = line
                tf.paragraphs[0].runs[0].font.name = MONO if mono else SANS
                started = True
            else:
                _para(tf, line, size=size, font=MONO if mono else SANS,
                      color=PAPER if b.kind != "prose" else MUTED,
                      spacing=1.35, space_before=min(14, size // 2))
    _notes(slide, s)
    _footer(slide, s)
    return slide


def draw_placeholder(prs, s: Slide):
    """The on-screen content is a screenshot or a view that is not a file yet."""
    slide = _blank(prs)
    _rule(slide, MARGIN, Inches(0.75), Inches(0.85))
    _text(slide, MARGIN, Inches(1.1), Inches(10.5), Inches(0.6), size=32, bold=True)
    slide.shapes[-1].text_frame.paragraphs[0].runs[0].text = s.title

    top = Inches(2.15)
    height = H - top - Inches(0.95)
    _panel(slide, MARGIN, top, W - 2 * MARGIN, height, color=INK, outline=AMBER)
    tf = _text(slide, MARGIN + Inches(0.6), top + Inches(0.5),
               W - 2 * MARGIN - Inches(1.2), height - Inches(1.0),
               size=13, color=AMBER, bold=True, anchor=MSO_ANCHOR.MIDDLE,
               align=PP_ALIGN.CENTER)
    tf.paragraphs[0].runs[0].text = "PUT ON THIS SLIDE"
    _para(tf, s.onscreen_note or "see deck.md", size=19, color=PAPER,
          align=PP_ALIGN.CENTER, spacing=1.4, space_before=18)
    for b in s.blocks:
        for line in b.lines:
            _para(tf, line, size=15, color=MUTED, align=PP_ALIGN.CENTER,
                  spacing=1.35, space_before=10)
    _notes(slide, s)
    _footer(slide, s)
    return slide


def build(slides: list[Slide]) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    clip_no = 0
    for s in slides:
        if s.onscreen_note == "__ACT__":
            draw_act(prs, s)
        elif s.number == "0":
            draw_title(prs, s)
        elif s.is_video:
            clip_no += 1
            draw_video(prs, s, clip_no)
        elif s.diagram in DIAGRAMS:
            DIAGRAMS[s.diagram](prs, s)
        else:
            code = next((b for b in s.blocks if b.kind == "code"), None)
            quote = next((b for b in s.blocks if b.kind == "quote"), None)
            listy = [b for b in s.blocks if b.kind in ("bullets", "prose")]
            if code:
                draw_code(prs, s, code)
            elif quote:
                draw_statement(prs, s, quote)
            elif listy:
                draw_bullets(prs, s, listy)
            else:
                draw_placeholder(prs, s)
    return prs


def audit(prs: Presentation) -> list[str]:
    """Report text running off a slide, since PowerPoint will not."""
    problems = []
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.left is None:
                continue
            if (shape.left < 0 or shape.top < 0
                    or shape.left + shape.width > prs.slide_width
                    or shape.top + shape.height > prs.slide_height):
                problems.append(f"slide {i}: a shape sits off the canvas")
            if not shape.has_text_frame:
                continue
            # Measure each paragraph at its own size. Taking the largest size on the
            # shape and applying it to every line reported the title slide, where a 66pt
            # heading sits above a 15pt link, as four inches over budget.
            width_in = shape.width / 914400
            used, first = 0.0, ""
            for para in shape.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs)
                if not text:
                    continue
                first = first or text
                size = max([int(r.font.size.pt) for r in para.runs if r.font.size] or [18])
                wrapped = max(1, int(len(text) * 0.0075 * size / max(width_in, 0.1)) + 1)
                used += wrapped * size * 1.3 / 72
            if first and used > shape.height / 914400 + 0.15:
                problems.append(f"slide {i}: {round(used, 1)}in of text in a "
                                f"{round(shape.height / 914400, 1)}in box - {first[:40]!r}")
    return problems


def main() -> int:
    if not DECK.exists():
        print(f"no deck at {DECK}", file=sys.stderr)
        return 1
    slides = parse(DECK.read_text(encoding="utf-8"))
    if not slides:
        print("parsed no slides - has the deck's heading format changed?", file=sys.stderr)
        return 1
    prs = build(slides)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(OUT)
    except PermissionError:
        print(f"{OUT.name} is open in PowerPoint - close it and run this again.",
              file=sys.stderr)
        return 2

    content = [s for s in slides if s.onscreen_note != "__ACT__"]
    videos = [s for s in content if s.is_video]
    missing = [f"clip-{i}" for i, s in enumerate(videos, 1)
               if not (CLIPS / f"clip-{i}.mp4").exists()]
    print(f"{OUT.relative_to(ROOT)}")
    print(f"  {len(slides)} slides ({len(content)} content, "
          f"{len(slides) - len(content)} act dividers)")
    print(f"  {len(videos)} video slides, {len(missing)} still placeholders: "
          f"{', '.join(missing) or 'none'}")
    empty = [s.number for s in content if not s.script]
    if empty:
        print(f"  no speaker notes on: {', '.join(empty)}")
    for problem in audit(prs):
        print(f"  ! {problem}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
